#!/usr/bin/env python3
"""
Build 10 IEEE-format research papers reusing the exact styles of
  VulnLens-AI_Research_Paper.docx  (A4, 2-column, Times New Roman,
  auto-numbered Heading1=I/II, Heading2=A/B, references=[1], tablehead=TABLE I).

Usage:  python build_papers.py
Output: C:/Users/YASH/Downloads/MINOR PROJECT/Papers/Paper_01_<slug>.docx  ...
"""
import copy
import importlib
import os
import re
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = r"C:/Users/YASH/Downloads/MINOR PROJECT/VulnLens-AI_Research_Paper.docx"
OUTDIR = r"C:/Users/YASH/Downloads/MINOR PROJECT/Papers"

AUTHORS = [
    ["Yash M. Panchal",
     "Dept. of Computer Science Engineering in Cyber Security",
     "Silver Oak University",
     "Ahmedabad, Gujarat, India"],
    ["Umang K. Nagwadiya",
     "Dept. of Computer Science Engineering in Cyber Security",
     "Silver Oak University",
     "Ahmedabad, Gujarat, India"],
]

# --------------------------------------------------------------------------
# low-level helpers
# --------------------------------------------------------------------------

def clear_body(doc):
    """Remove every body child except the trailing sectPr."""
    body = doc.element.body
    sectPr = body.find(qn('w:sectPr'))
    for child in list(body):
        if child is not sectPr:
            body.remove(child)


def _mk(parent, tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn('w:' + k), str(v))
    parent.append(el)
    return el


def add_para(doc, style, text="", italic_spans=()):
    """Add a paragraph with the given style.  _underscored_ spans become italic."""
    p = doc.add_paragraph(style=style)
    if text:
        parts = re.split(r'(_[^_]+_)', text)
        for part in parts:
            if not part:
                continue
            if part.startswith('_') and part.endswith('_') and len(part) > 2:
                r = p.add_run(part[1:-1])
                r.italic = True
            else:
                p.add_run(part)
    return p


def add_author_table(doc):
    """1x2 borderless author block, cloned from the template's structure."""
    tbl = OxmlElement('w:tbl')
    tblPr = _mk(tbl, 'w:tblPr')
    _mk(tblPr, 'w:tblW', w=0, type='dxa')
    _mk(tblPr, 'w:jc', val='center')
    cellMar = _mk(tblPr, 'w:tblCellMar')
    _mk(cellMar, 'w:left', w=57, type='dxa')
    _mk(cellMar, 'w:right', w=57, type='dxa')
    _mk(tblPr, 'w:tblLook', val='0000', firstRow=0, lastRow=0,
        firstColumn=0, lastColumn=0, noHBand=0, noVBand=0)

    grid = _mk(tbl, 'w:tblGrid')
    for _ in range(2):
        _mk(grid, 'w:gridCol', w=1232)

    tr = _mk(tbl, 'w:tr')
    trPr = _mk(tr, 'w:trPr')
    _mk(trPr, 'w:jc', val='center')

    for lines in AUTHORS:
        tc = _mk(tr, 'w:tc')
        tcPr = _mk(tc, 'w:tcPr')
        _mk(tcPr, 'w:tcW', w=0, type='dxa')
        for i, line in enumerate(lines):
            p = OxmlElement('w:p')
            pPr = _mk(p, 'w:pPr')
            _mk(pPr, 'w:pStyle', val='Author')
            if i == len(lines) - 1:
                _mk(pPr, 'w:spacing', after=0)
            r = _mk(p, 'w:r')
            if i == 0:
                _mk(r, 'w:rPr')
            t = _mk(r, 'w:t')
            t.set(qn('xml:space'), 'preserve')
            t.text = line
            tc.append(p)

    doc.element.body.sectPr.addprevious(tbl)
    return tbl


def add_data_table(doc, spec):
    """Bordered data table: spec = {'cols':n,'rows':[[...]],'head':...,'footnote':...}"""
    rows = spec["rows"]
    ncols = spec["cols"]
    total = 4600
    each = total // ncols

    tbl = OxmlElement('w:tbl')
    tblPr = _mk(tbl, 'w:tblPr')
    _mk(tblPr, 'w:tblW', w=0, type='dxa')
    _mk(tblPr, 'w:jc', val='center')
    borders = _mk(tblPr, 'w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        _mk(borders, 'w:' + edge, val='single', sz=4, space=0, color='000000')
    cellMar = _mk(tblPr, 'w:tblCellMar')
    _mk(cellMar, 'w:left', w=57, type='dxa')
    _mk(cellMar, 'w:right', w=57, type='dxa')
    _mk(tblPr, 'w:tblLook', val='0000', firstRow=0, lastRow=0,
        firstColumn=0, lastColumn=0, noHBand=0, noVBand=0)

    grid = _mk(tbl, 'w:tblGrid')
    for _ in range(ncols):
        _mk(grid, 'w:gridCol', w=each)

    for ri, row in enumerate(rows):
        tr = _mk(tbl, 'w:tr')
        trPr = _mk(tr, 'w:trPr')
        _mk(trPr, 'w:jc', val='center')
        for ci in range(ncols):
            txt = row[ci] if ci < len(row) else ""
            tc = _mk(tr, 'w:tc')
            tcPr = _mk(tc, 'w:tcPr')
            _mk(tcPr, 'w:tcW', w=0, type='dxa')
            if ri == 0:
                _mk(tcPr, 'w:shd', val='clear', color='auto', fill='F2F2F2')
            p = OxmlElement('w:p')
            pPr = _mk(p, 'w:pPr')
            _mk(pPr, 'w:spacing', before=0, after=0, line=200, lineRule='auto')
            _mk(pPr, 'w:jc', val='center')
            r = _mk(p, 'w:r')
            rPr = _mk(r, 'w:rPr')
            _mk(rPr, 'w:sz', val=14)
            _mk(rPr, 'w:szCs', val=14)
            if ri == 0:
                _mk(rPr, 'w:b')
                _mk(rPr, 'w:bCs')
            t = _mk(r, 'w:t')
            t.set(qn('xml:space'), 'preserve')
            t.text = txt
            tc.append(p)
            tr.append(tc)

    doc.element.body.sectPr.addprevious(tbl)

    if spec.get("head"):
        add_para(doc, 'tablehead', spec["head"])
    if spec.get("footnote"):
        add_para(doc, 'tablefootnote', spec["footnote"])


# --------------------------------------------------------------------------
# document assembly
# --------------------------------------------------------------------------

def build_paper(data, out_path):
    doc = Document(TEMPLATE)
    clear_body(doc)

    add_para(doc, 'papertitle', data["title"])
    add_author_table(doc)
    add_para(doc, 'Abstract', data["abstract"])
    add_para(doc, 'Keywords', "Keywords\u2014" + data["keywords"])

    # index tables by the section they belong after
    tables_by_section = {}
    for t in data.get("tables", []):
        tables_by_section.setdefault(t.get("after_section"), []).append(t)

    for heading, blocks in data["sections"]:
        add_para(doc, 'Heading 1', heading)
        for block in blocks:
            if isinstance(block, tuple):
                sub, paras = block
                if sub:
                    add_para(doc, 'Heading 2', sub)
                for para in paras:
                    add_para(doc, 'Body Text', para)
            elif isinstance(block, dict) and "rows" in block:
                add_data_table(doc, block)
            else:
                add_para(doc, 'Body Text', block)
        for t in tables_by_section.get(heading, []):
            add_data_table(doc, t)

    # any table whose anchor section was not found -> append at the end
    placed = set()
    for heading, _ in data["sections"]:
        placed.update(id(t) for t in tables_by_section.get(heading, []))
    for t in data.get("tables", []):
        if id(t) not in placed:
            add_data_table(doc, t)

    add_para(doc, 'Heading 5', 'Acknowledgment')
    add_para(doc, 'Body Text', data["ack"])

    add_para(doc, 'Heading 5', 'References')
    for ref in data["refs"]:
        add_para(doc, 'references', ref)

    doc.save(out_path)
    return out_path


def main():
    os.makedirs(OUTDIR, exist_ok=True)
    built = []
    for i in range(1, 11):
        mod = importlib.import_module("p%d" % i)
        data = mod.PAPER
        name = "Paper_%02d_%s.docx" % (data["num"], data["slug"])
        path = os.path.join(OUTDIR, name)
        build_paper(data, path)
        built.append(path)
        print("built  %-58s %8d bytes" % (name, os.path.getsize(path)))
    print("\n%d papers written to %s" % (len(built), OUTDIR))
    return 0


if __name__ == "__main__":
    sys.exit(main())
